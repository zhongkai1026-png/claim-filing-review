"""台账 Excel 导入服务：以系统导出的台账 Excel 为模板，按表头映射同步到系统台账。

- 表头按系统导出台账格式（LEDGER_COLUMNS）映射，并支持常见别名
- 债权编号：文件中有编号且项目内已存在 → 更新该条；否则自动分配新编号（PC-XXX）
- 金额：支持千分位/人民币符号/「万」单位
- 状态：待登记/已登记/已复核 自动映射
- 返回导入统计与失败明细
"""
import re
from io import BytesIO

from openpyxl import load_workbook

from ..models import Claim

# 系统导出台账表头 -> 字段
HEADER_MAP = {
    "债权编号": "claim_no",
    "债权人名称": "name",
    "统一社会信用代码/身份证号码": "credit_code",
    "法定代表人/负责人": "legal_rep",
    "联系方式": "contact_info",
    "联系地址": "address",
    "代理人姓名": "agent_name",
    "代理人身份证号": "agent_id_no",
    "代理人联系方式": "agent_contact",
    "代理人地址": "agent_address",
    "开户名称": "account_name",
    "开户行": "bank_name",
    "账号": "account_no",
    "债权种类": "claim_category",
    "涉诉涉执情况": "litigation_status",
    "申报债权性质": "claim_nature",
    "申报债权时间": "claim_date",
    "申报本金": "principal",
    "申报利息": "interest",
    "申报其他": "others",
    "申报总额": "claim_amount",
    "审核本金": "audited_principal",
    "审核利息": "audited_interest",
    "审核其他": "audited_others",
    "审核总额": "audited_amount",
    "债权性质": "debt_nature",
    "担保情况备注": "guarantee_notes",
    "债权类型": "claim_type",
    "状态": "status",
}

# 常见别名（外部模板兼容）
ALIASES = {
    "债权人": "name", "债权人名称": "name", "单位名称": "name",
    "申报人": "name", "债权人姓名": "name",
    "信用代码": "credit_code", "统一社会信用代码": "credit_code",
    "身份证号": "credit_code", "身份证号码": "credit_code",
    "法定代表人": "legal_rep", "法人代表": "legal_rep", "负责人": "legal_rep",
    "电话": "contact_info", "联系电话": "contact_info", "手机": "contact_info",
    "联系地址": "address", "地址": "address", "通讯地址": "address",
    "申报金额": "claim_amount", "申报总额": "claim_amount", "合计金额": "claim_amount",
    "本金": "principal", "申报本金": "principal",
    "利息": "interest", "申报利息": "interest",
    "其他": "others", "申报其他": "others", "其他费用": "others",
    "审核金额": "audited_amount", "确认金额": "audited_amount",
    "债权种类": "claim_category", "债权类别": "claim_category",
    "申报债权性质": "claim_nature", "债权性质": "claim_nature", "性质": "claim_nature",
    "申报时间": "claim_date", "债权发生时间": "claim_date",
    "开户行": "bank_name", "开户银行": "bank_name",
    "账号": "account_no", "银行账号": "account_no", "开户账号": "account_no",
    "担保情况": "guarantee_notes", "担保备注": "guarantee_notes",
    "债权类型": "claim_type", "类型": "claim_type",
    "涉诉": "litigation_status", "涉诉情况": "litigation_status",
}

# 导入时忽略的列（来源材料/创建时间 由系统维护）
SKIP_FIELDS = {"source_files", "created_at"}
MONEY_FIELDS = {"principal", "interest", "others", "claim_amount",
                "audited_principal", "audited_interest", "audited_others",
                "audited_amount"}
STATUS_MAP = {
    "待登记": "pending", "已登记": "registered", "已复核": "reviewed",
    "pending": "pending", "registered": "registered", "reviewed": "reviewed",
    "": "pending",
}


def parse_amount(v) -> float:
    """金额解析：数字 / 千分位 / 人民币符号 / 「万」单位 / 空 -> 0。"""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).strip().replace(",", "").replace("￥", "").replace("¥", "").replace("元", "")
    if not s:
        return 0.0
    m = re.match(r"^([\d.]+)(万)?$", s)
    if m:
        amt = float(m.group(1))
        if m.group(2):
            amt *= 10000
        return round(amt, 2)
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def _norm_header(s) -> str:
    """表头规范化：去掉所有空白/换行/首尾符号，避免映射失败。"""
    return re.sub(r"\s+", "", str(s or "")).strip().strip("＊*·．. ")


# 规范化后的映射表（匹配时忽略表头里的空格/换行）
_NORM_HEADER = {_norm_header(k): v for k, v in HEADER_MAP.items()}
_NORM_ALIAS = {_norm_header(k): v for k, v in ALIASES.items()}

# 文本类字段（长数字如信用代码/账号/身份证需按文本处理，避免精度丢失）
TEXT_FIELDS = {"name", "credit_code", "legal_rep", "contact_info", "address",
               "agent_name", "agent_id_no", "agent_contact", "agent_address",
               "account_name", "bank_name", "account_no", "claim_category",
               "litigation_status", "claim_nature", "claim_date", "debt_nature",
               "guarantee_notes", "claim_type"}


def _to_text(v) -> str:
    """文本字段取值：数字（长整数/科学计数法）转为完整字符串。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return format(v, ".0f") if abs(v) > 1e12 else str(v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def build_header_map(header_row: tuple) -> dict:
    """表头行 -> {列索引: 字段名}（含别名与规范化，SKIP 字段跳过）。"""
    mapping = {}
    for idx, h in enumerate(header_row):
        key = _norm_header(h)
        if not key:
            continue
        field = _NORM_HEADER.get(key) or _NORM_ALIAS.get(key)
        if field and field not in SKIP_FIELDS:
            mapping[idx] = field
    return mapping


def import_ledger_excel(data: bytes, project_id: int, db) -> dict:
    """执行导入，返回统计与错误明细。"""
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception as e:  # noqa: BLE001
        return {"success": 0, "created": 0, "updated": 0, "failed": 0,
                "message": f"文件解析失败：{e}", "errors": []}

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"success": 0, "created": 0, "updated": 0, "failed": 0,
                "message": "文件为空", "errors": []}

    header = rows[0]
    mapping = build_header_map(header)
    if not mapping:
        return {"success": 0, "created": 0, "updated": 0, "failed": 0,
                "message": "无法识别表头，请使用系统导出的台账模板（含「债权人名称」等列）",
                "errors": []}

    # 项目内已有债权索引（按编号）
    existing = {c.claim_no: c for c in db.query(Claim).filter(
        Claim.project_id == project_id).all()}
    # 下一个可用序号
    nums = []
    for no in existing:
        m = re.search(r"(\d+)$", no or "")
        if m:
            nums.append(int(m.group(1)))
    next_seq = (max(nums) + 1) if nums else 1

    created = updated = failed = 0
    errors = []
    for ridx, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue  # 空行

        data_dict = {}
        for idx, field in mapping.items():
            if idx < len(row) and row[idx] is not None:
                val = row[idx]
                # 文本类字段：长数字（信用代码/账号/身份证）按文本还原，避免精度丢失
                if field in TEXT_FIELDS:
                    val = _to_text(val)
                data_dict[field] = val

        name = str(data_dict.get("name") or "").strip()
        if not name:
            failed += 1
            errors.append({"row": ridx, "reason": "债权人名称为空，无法导入"})
            continue

        # 金额解析
        for f in MONEY_FIELDS:
            if f in data_dict:
                data_dict[f] = parse_amount(data_dict[f])
        # 日期统一为「2025年4月5日」
        if data_dict.get("claim_date"):
            from .formatter import format_date_cn
            data_dict["claim_date"] = format_date_cn(data_dict["claim_date"])
        # 状态
        if data_dict.get("status") not in (None, ""):
            st = str(data_dict["status"]).strip()
            data_dict["status"] = STATUS_MAP.get(st, st)

        # 定位记录：按债权编号更新，否则新建
        claim_no = str(data_dict.get("claim_no") or "").strip()
        claim = existing.get(claim_no)
        if claim is None:
            # 导入台账的行没有上传材料，解析状态标记为 empty（无材料无需解析），
            # 避免长期停留在 pending 造成"待解析"误解，或被解析自愈误判为失败
            claim = Claim(project_id=project_id,
                          claim_no=f"PC-{next_seq:03d}",
                          parse_status="empty")
            db.add(claim)
            next_seq += 1
            created += 1
        else:
            updated += 1

        for field, val in data_dict.items():
            if field in ("claim_no",) or field in SKIP_FIELDS:
                continue
            setattr(claim, field, val)
        # 申报总额未填时按分项自动求和
        if not (claim.claim_amount or 0):
            claim.claim_amount = round(
                float(claim.principal or 0) + float(claim.interest or 0)
                + float(claim.others or 0), 2)
        # 审核总额未填时按分项自动求和
        if not (claim.audited_amount or 0):
            claim.audited_amount = round(
                float(claim.audited_principal or 0)
                + float(claim.audited_interest or 0)
                + float(claim.audited_others or 0), 2)
        db.flush()
        existing[claim.claim_no] = claim  # 同文件内后续行引用

    db.commit()
    total = created + updated
    msg = (f"导入完成：成功 {total} 条（新增 {created} / 更新 {updated}）"
           + (f"，失败 {failed} 条" if failed else ""))
    # 表头识别诊断（按列序）
    mapped_cols = []
    for idx in sorted(mapping):
        mapped_cols.append(f"{idx + 1}列→{mapping[idx]}")
    return {"success": total, "created": created, "updated": updated,
            "failed": failed, "message": msg, "errors": errors[:50],
            "mapped_columns": mapped_cols}
